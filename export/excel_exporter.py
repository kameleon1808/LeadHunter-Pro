"""
ExcelExporter — exports campaign leads to a formatted Excel workbook.

Output file contains three sheets:
  Sheet 1 "Leads"       — all lead records with formatted columns
  Sheet 2 "Stats"       — campaign statistics summary
  Sheet 3 "Failed URLs" — URLs that failed during enrichment

Requires: openpyxl>=3.1.0
"""

import os
import re
from datetime import date, datetime
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import DatabaseManager
from database.models import URL

# ── Colour palette ──────────────────────────────────────────────────────────
_HEADER_BG    = "1F4E79"   # dark navy
_HEADER_FG    = "FFFFFF"   # white
_ROW_ODD_BG   = "FFFFFF"   # white
_ROW_EVEN_BG  = "EBF3FB"   # pale blue
_STATS_KEY_BG = "D6E4F0"   # light steel blue
_SHEET_TAB_LEADS  = "2E75B6"
_SHEET_TAB_STATS  = "375623"
_SHEET_TAB_FAILED = "C00000"

# ── Leads sheet column definitions (header, model attr, max width) ──────────
_LEAD_COLUMNS = [
    ("#",             None,              5),
    ("Company Name",  "company_name",   30),
    ("Industry",      "industry",       20),
    ("Size",          "company_size",   15),
    ("Description",   "description",    45),
    ("Contact Name",  "contact_name",   22),
    ("Contact Title", "contact_title",  20),
    ("Contact Email", "contact_email",  30),
    ("Contact LinkedIn", "contact_linkedin", 35),
    ("Company Email", "company_email",  30),
    ("Company Phone", "company_phone",  18),
    ("Address",       "address",        30),
    ("Website",       "website_url",    30),
    ("Clients Info",  "clients_info",   40),
    ("Quality Score", "quality_score",  14),
]


class ExcelExporter:
    """Exports campaign data to a formatted .xlsx file."""

    def __init__(self, db_manager: DatabaseManager) -> None:
        self._db = db_manager

    # ── Public API ───────────────────────────────────────────────────────────

    def export_campaign(self, campaign_id: int, output_path: str) -> str:
        """
        Create a formatted Excel workbook for *campaign_id*.

        *output_path* may be a directory (file is named automatically) or a
        full file path.  Returns the absolute path of the created file.
        """
        campaign = self._db.get_campaign(campaign_id)
        if campaign is None:
            raise ValueError(f"Campaign {campaign_id} not found")

        leads      = self._db.get_leads(campaign_id)
        stats      = self._db.get_campaign_stats(campaign_id)
        failed_urls = self._get_failed_urls(campaign_id)

        wb = Workbook()

        # Sheet 1 — Leads
        ws_leads = wb.active
        ws_leads.title = "Leads"
        ws_leads.sheet_properties.tabColor = _SHEET_TAB_LEADS
        self._format_leads_sheet(ws_leads, leads)

        # Sheet 2 — Stats
        ws_stats = wb.create_sheet("Stats")
        ws_stats.sheet_properties.tabColor = _SHEET_TAB_STATS
        self._format_stats_sheet(ws_stats, stats, campaign)

        # Sheet 3 — Failed URLs
        ws_failed = wb.create_sheet("Failed URLs")
        ws_failed.sheet_properties.tabColor = _SHEET_TAB_FAILED
        self._format_failed_sheet(ws_failed, failed_urls)

        file_path = self._resolve_output_path(output_path, campaign)
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        wb.save(file_path)
        return os.path.abspath(file_path)

    # ── Sheet formatters ─────────────────────────────────────────────────────

    def _format_leads_sheet(self, ws, leads: list) -> None:
        """Write and format the Leads sheet."""
        headers = [col[0] for col in _LEAD_COLUMNS]
        self._write_header_row(ws, headers)

        for row_idx, lead in enumerate(leads, start=2):
            is_even = (row_idx % 2 == 0)
            bg = _ROW_EVEN_BG if is_even else _ROW_ODD_BG
            fill = PatternFill("solid", fgColor=bg)

            for col_idx, (_, attr, _) in enumerate(_LEAD_COLUMNS, start=1):
                if attr is None:
                    value = row_idx - 1        # row number column
                else:
                    value = getattr(lead, attr, None)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = fill
                cell.alignment = Alignment(wrap_text=False, vertical="top")

        # Filters on header row
        ws.auto_filter.ref = ws.dimensions

        # Freeze top row
        ws.freeze_panes = "A2"

        # Auto-width (capped per column definition)
        for col_idx, (header, attr, max_w) in enumerate(_LEAD_COLUMNS, start=1):
            letter = get_column_letter(col_idx)
            col_width = len(header)
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value is not None:
                        col_width = max(col_width, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(col_width + 2, max_w)

    def _format_stats_sheet(self, ws, stats: dict, campaign) -> None:
        """Write and format the Stats summary sheet."""
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 30

        title_font  = Font(bold=True, size=14, color=_HEADER_FG)
        title_fill  = PatternFill("solid", fgColor=_HEADER_BG)
        key_fill    = PatternFill("solid", fgColor=_STATS_KEY_BG)
        key_font    = Font(bold=True)
        value_align = Alignment(horizontal="left")

        # Title row
        ws.merge_cells("A1:B1")
        title_cell = ws["A1"]
        title_cell.value = "Campaign Statistics"
        title_cell.font  = title_font
        title_cell.fill  = title_fill
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        total    = stats["total_urls"]
        done     = stats["completed_urls"] + stats["failed_urls"]
        success  = stats["completed_urls"]
        failed   = stats["failed_urls"]
        leads    = stats["total_leads"]
        lead_pct = f"{leads / success * 100:.1f}%" if success else "N/A"
        succ_pct = f"{success / total * 100:.1f}%" if total else "N/A"

        rows = [
            ("Campaign Name",    campaign.name),
            ("Niche",            campaign.niche),
            ("Target Country",   campaign.target_country or "—"),
            ("Status",           campaign.status.capitalize()),
            ("Created At",       campaign.created_at.strftime("%Y-%m-%d %H:%M") if campaign.created_at else "—"),
            ("",                 ""),
            ("Total URLs",       total),
            ("Processed",        done),
            ("Completed",        success),
            ("Failed",           failed),
            ("Success Rate",     succ_pct),
            ("",                 ""),
            ("Leads Found",      leads),
            ("Lead Conversion",  lead_pct),
        ]

        for r_idx, (key, value) in enumerate(rows, start=2):
            ws.row_dimensions[r_idx].height = 18
            key_cell   = ws.cell(row=r_idx, column=1, value=key)
            value_cell = ws.cell(row=r_idx, column=2, value=value)
            if key:
                key_cell.fill  = key_fill
                key_cell.font  = key_font
            value_cell.alignment = value_align

    def _format_failed_sheet(self, ws, failed_urls: list) -> None:
        """Write and format the Failed URLs sheet."""
        headers = ["#", "URL", "Error Message", "Retry Count"]
        col_widths = [5, 55, 50, 12]
        self._write_header_row(ws, headers)

        for col_idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        for row_idx, url_rec in enumerate(failed_urls, start=2):
            is_even = (row_idx % 2 == 0)
            bg   = _ROW_EVEN_BG if is_even else _ROW_ODD_BG
            fill = PatternFill("solid", fgColor=bg)
            values = [
                row_idx - 1,
                url_rec.url,
                url_rec.error_message or "",
                url_rec.retry_count,
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = fill
                cell.alignment = Alignment(wrap_text=False, vertical="top")

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"

    # ── Helpers ──────────────────────────────────────────────────────────────

    def _write_header_row(self, ws, headers: List[str]) -> None:
        """Write and style the header row."""
        header_font = Font(bold=True, color=_HEADER_FG, size=11)
        header_fill = PatternFill("solid", fgColor=_HEADER_BG)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
        ws.row_dimensions[1].height = 20

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = header_align

    def _get_failed_urls(self, campaign_id: int) -> list:
        """Return URL records with status 'failed' for the campaign."""
        from sqlalchemy.orm import Session
        with Session(self._db.engine) as session:
            records = (
                session.query(URL)
                .filter(URL.campaign_id == campaign_id, URL.status == "failed")
                .order_by(URL.processed_at)
                .all()
            )
            for r in records:
                session.expunge(r)
            return records

    @staticmethod
    def _resolve_output_path(output_path: str, campaign) -> str:
        """
        Resolve output_path to a full file path.

        If *output_path* has an .xlsx extension it is used as-is (explicit file).
        Otherwise it is treated as a directory and the filename is auto-generated:
            leads_{safe_name}_campaign{id}_{date}.xlsx
        """
        _, ext = os.path.splitext(output_path)
        if ext.lower() in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            return output_path

        safe_name = re.sub(r"[^\w\s-]", "", campaign.name.lower())
        safe_name = re.sub(r"[\s-]+", "_", safe_name).strip("_")
        date_str  = date.today().strftime("%Y-%m-%d")
        filename  = f"leads_{safe_name}_campaign{campaign.id}_{date_str}.xlsx"
        return os.path.join(output_path, filename)
